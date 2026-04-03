"""Work Email Tool (IMAP) — อ่านเมลองค์กรและสรุปพร้อมไฟล์แนบ"""

import asyncio
import email
import imaplib
import ssl
import re
from datetime import datetime, timedelta
from email.header import decode_header
from email.utils import parseaddr
import io
import urllib.parse
from dataclasses import dataclass, field
import html
import hashlib

# extractors will be lazily imported

from tools.base import BaseTool
from core.api_keys import get_api_key
from core.config import (
    WORK_IMAP_PORT,
    WORK_EMAIL_MAX_RESULTS,
    WORK_EMAIL_ATTACHMENT_MAX_MB,
)
from core import db
from core.llm import llm_router
from core.user_manager import get_user_by_id, get_preference
from core.logger import get_logger

log = get_logger(__name__)


@dataclass
class ParsedArgs:
    time_range: str = "1d"
    time_label: str = "วันนี้"
    force: bool = False
    filters: dict = field(default_factory=dict)
    search_text: str = ""


@dataclass
class AttachmentData:
    filename: str
    size_bytes: int
    mime_type: str
    content: str | None = None
    status: str = "extracted"


@dataclass
class EmailData:
    message_id: str
    subject: str
    sender: str
    date: str
    body: str
    attachments: list[AttachmentData] = field(default_factory=list)


class WorkEmailTool(BaseTool):
    name = "work_email"
    description = "สรุปอีเมลที่ทำงาน (IMAP) พร้อมอ่านไฟล์แนบ ค้นหาได้"
    commands = ["/wm", "/workmail"]
    preferred_tier = "mid"

    TIME_RANGES = {
        "today": ("1d", "วันนี้"),
        "1d": ("1d", "วันนี้"),
        "3d": ("3d", "3 วันล่าสุด"),
        "7d": ("7d", "7 วันล่าสุด"),
        "14d": ("14d", "14 วันล่าสุด"),
        "30d": ("30d", "30 วันล่าสุด"),
    }

    THAI_TIME_PATTERNS = (
        re.compile(r"(?:(?:ใน|ย้อนหลัง)\s*)?(\d{1,3})\s*วัน(?:ที่ผ่านมา|ล่าสุด)?"),
    )

    def _extract_time_range(self, text: str) -> tuple[str, str, str]:
        time_range = "1d"
        time_label = "วันนี้"
        remaining = text

        for pattern in self.THAI_TIME_PATTERNS:
            match = pattern.search(remaining)
            if not match:
                continue
            days = match.group(1)
            time_range = f"{days}d"
            time_label = "วันนี้" if days == "1" else f"{days} วันล่าสุด"
            remaining = (remaining[:match.start()] + " " + remaining[match.end():]).strip()
            return time_range, time_label, remaining

        return time_range, time_label, remaining

    def _parse_args(self, args: str) -> ParsedArgs:
        """แยก time_range, force, filters, search_text"""
        parsed = ParsedArgs()
        if not args:
            return parsed

        original_args = args.strip()
        normalized_args = re.sub(r"\bforce\b", " ", original_args, flags=re.IGNORECASE).strip()

        parsed.force = bool(re.search(r"\bforce\b", original_args, flags=re.IGNORECASE))

        parsed.time_range, parsed.time_label, remaining_args = self._extract_time_range(normalized_args)

        original_tokens = remaining_args.split()
        search_tokens = []

        for token in original_tokens:
            token_lower = token.lower()
            if token_lower in self.TIME_RANGES:
                parsed.time_range, parsed.time_label = self.TIME_RANGES[token_lower]
            elif re.match(r"^\d+d$", token_lower):
                parsed.time_range = token_lower
                parsed.time_label = f"{token_lower[:-1]} วันล่าสุด"
            elif ":" in token_lower and not token.isspace() and not token.startswith("http"):
                key, val = token.split(":", 1)
                key_lower = key.lower()
                if key_lower in ["from", "to", "subject", "body", "folder"]:
                    parsed.filters[key_lower] = val
                else:
                    search_tokens.append(token)
            else:
                search_tokens.append(token)

        parsed.search_text = " ".join(search_tokens)
        return parsed

    def _resolve_imap_credentials(self, user_id: str) -> tuple[str, str, str] | None:
        host = get_api_key(user_id, "work_imap_host")
        username = get_api_key(user_id, "work_imap_user")
        password = get_api_key(user_id, "work_imap_password")

        if not host or not username or not password:
            return None

        return host, username, password

    def _connect_imap(self, host: str, username: str, password: str) -> imaplib.IMAP4_SSL:
        try:
            conn = imaplib.IMAP4_SSL(host, WORK_IMAP_PORT)
        except ssl.SSLCertVerificationError:
            ctx = ssl.create_default_context()
            ctx.check_hostname = False
            ctx.verify_mode = ssl.CERT_NONE
            conn = imaplib.IMAP4_SSL(host, WORK_IMAP_PORT, ssl_context=ctx)
            log.warning("IMAP SSL: certificate verification disabled")
            
        conn.login(username, password)
        return conn

    def _build_search_criteria(self, parsed: ParsedArgs) -> str:
        """สร้าง IMAP SEARCH string"""
        criteria = []
        
        # ไม่ใช้ UNSEEN เพราะเมลองค์กรถูกอ่านผ่าน client อื่นอยู่แล้ว
        # ใช้ DB-level dedup (processed_emails) แทน
        # force = ข้ามการเช็ค processed_emails (สรุปซ้ำได้)
        # Parse time_range e.g. "7d"
        days = 1
        m = re.match(r"^(\d+)d$", parsed.time_range)
        if m:
            days = int(m.group(1))
            
        since_date = (datetime.now() - timedelta(days=days)).strftime("%d-%b-%Y")
        criteria.append(f'SINCE {since_date}')
        
        # For better Thai compatibility, we do not push TEXT or SUBJECT search to IMAP server 
        # unless it is English only, but to be simple and robust, we fetch everything SINCE and filter in python.
        # We only send SINCE to the server.
        pass
            
        if not criteria:
            criteria.append("ALL")
            
        return f"({' '.join(criteria)})"
        
    def _decode_header(self, raw: str) -> str:
        if not raw:
            return ""
        try:
            parts = decode_header(raw)
            decoded = []
            for data, charset in parts:
                if isinstance(data, bytes):
                    charset = charset or "utf-8"
                    for enc in [charset, "utf-8", "tis-620", "windows-874"]:
                        try:
                            decoded.append(data.decode(enc))
                            break
                        except (UnicodeDecodeError, LookupError):
                            continue
                    else:
                        decoded.append(data.decode("utf-8", errors="replace"))
                else:
                    decoded.append(str(data))
            return " ".join(decoded).strip()
        except Exception as e:
            log.warning(f"Error decoding header {raw}: {e}")
            return str(raw)

    def _extract_body(self, msg: email.message.Message) -> str:
        text_parts = []
        html_parts = []
        if msg.is_multipart():
            for part in msg.walk():
                content_type = part.get_content_type()
                content_disposition = str(part.get("Content-Disposition"))

                if "attachment" not in content_disposition:
                    if content_type == "text/plain":
                        payload = part.get_payload(decode=True)
                        if payload:
                            charset = part.get_content_charset() or "utf-8"
                            text_parts.append(self._decode_bytes(payload, charset))
                    elif content_type == "text/html":
                        payload = part.get_payload(decode=True)
                        if payload:
                            charset = part.get_content_charset() or "utf-8"
                            html_str = self._decode_bytes(payload, charset)
                            try:
                                from bs4 import BeautifulSoup
                                soup = BeautifulSoup(html_str, "html.parser")
                                for br in soup.find_all("br"): br.replace_with("\n")
                                for p in soup.find_all("p"): p.insert_after("\n")
                                for tr in soup.find_all("tr"):
                                    for th in tr.find_all(["th", "td"]):
                                        th.insert_after(" | ")
                                    tr.insert_after("\n")
                                text = soup.get_text()
                            except ImportError:
                                text = re.sub(r"<br\s*/?>", "\n", html_str, flags=re.IGNORECASE)
                                text = re.sub(r"</p>", "\n\n", text, flags=re.IGNORECASE)
                                text = re.sub(r"</tr>", "\n", text, flags=re.IGNORECASE)
                                text = re.sub(r"</td>", " | ", text, flags=re.IGNORECASE)
                                text = re.sub(r"<[^>]+>", " ", text)
                                
                            html_parts.append(html.unescape(text).strip())
        else:
            payload = msg.get_payload(decode=True)
            if payload:
                charset = msg.get_content_charset() or "utf-8"
                text_parts.append(self._decode_bytes(payload, charset))
                
        # prefer text/plain if available, else combine what we have
        if text_parts:
            return "\n".join(text_parts).strip()
        return "\n".join(html_parts).strip()

    def _decode_bytes(self, data: bytes, charset: str) -> str:
        for enc in [charset, "utf-8", "tis-620", "windows-874"]:
            if not enc:
                continue
            try:
                return data.decode(enc)
            except (UnicodeDecodeError, LookupError):
                continue
        return data.decode("utf-8", errors="replace")

    def _extract_pdf(self, data: bytes) -> str:
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = ""
                for page in pdf.pages[:5]:  # limit to 5 pages
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
                return text[:3000]
        except Exception as e:
            log.warning(f"Failed to extract PDF: {e}")
            return ""

    def _extract_docx(self, data: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs])
            return text[:3000]
        except Exception as e:
            log.warning(f"Failed to extract DOCX: {e}")
            return ""

    def _extract_xlsx(self, data: bytes) -> str:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True)
            ws = wb.active
            text = ""
            for row in list(ws.rows)[:50]: # limit to 50 rows
                row_vals = [str(cell.value) if cell.value is not None else "" for cell in row]
                if any(row_vals):
                    text += " | ".join(row_vals) + "\n"
            return text[:3000]
        except Exception as e:
            log.warning(f"Failed to extract XLSX: {e}")
            return ""

    def _extract_text_file(self, data: bytes, charset: str) -> str:
        return self._decode_bytes(data, charset)[:3000]

    def _process_attachments(self, msg: email.message.Message) -> list[AttachmentData]:
        attachments = []
        for part in msg.walk():
            if part.get_content_maintype() == "multipart":
                continue
            
            content_disposition = str(part.get("Content-Disposition"))
            if "attachment" not in content_disposition and "inline" not in content_disposition:
                continue
                
            filename = part.get_filename()
            if not filename:
                continue
                
            filename = self._decode_header(urllib.parse.unquote(filename))
            payload = part.get_payload(decode=True)
            if not payload:
                continue

            size_bytes = len(payload)
            mime_type = part.get_content_type()
            
            att = AttachmentData(filename=filename, size_bytes=size_bytes, mime_type=mime_type)
            
            extracted_count = sum(1 for a in attachments if a.status == "extracted")
            if extracted_count >= 3:
                att.content = None
                att.status = "skipped_limit"
                attachments.append(att)
                if len(attachments) >= 5: # Limit reports
                    break
                continue
                
            if size_bytes > WORK_EMAIL_ATTACHMENT_MAX_MB * 1024 * 1024:
                att.content = None
                att.status = "too_large"
                attachments.append(att)
                continue
                
            content = None
            try:
                if mime_type == "application/pdf" or filename.lower().endswith(".pdf"):
                    content = self._extract_pdf(payload)
                elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.lower().endswith(".docx"):
                    content = self._extract_docx(payload)
                elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or filename.lower().endswith(".xlsx"):
                    content = self._extract_xlsx(payload)
                elif mime_type.startswith("text/") or filename.lower().endswith(".csv") or filename.lower().endswith(".txt"):
                    charset = part.get_content_charset() or "utf-8"
                    content = self._extract_text_file(payload, charset)
                elif mime_type.startswith("image/") or filename.lower().endswith((".png", ".jpg", ".jpeg")):
                    att.status = "image"
                else:
                    att.status = "unsupported"
            except Exception as e:
                log.error(f"Error extracting {filename}: {e}")
                att.status = "error"
                
            if content is not None:
                att.content = content.strip()
                if not att.content:
                    att.status = "empty_or_scanned"
                    att.content = None
                else:
                    att.status = "extracted"
            
            attachments.append(att)
            
        return attachments

    def _sync_fetch_all(self, user_id: str, parsed: ParsedArgs, imap_credentials: tuple[str, str, str]) -> tuple[list[EmailData], int]:
        skipped = 0
        emails_data = []
        conn = None
        try:
            conn = self._connect_imap(*imap_credentials)
            
            # Select Folder
            folder = parsed.filters.get('folder', 'INBOX')
            status, _ = conn.select(folder)
            if status != "OK":
                log.warning(f"Failed to select folder: {folder}")
                return [], 0
            
            criteria = self._build_search_criteria(parsed)
            log.info(f"IMAP SEARCH in {folder}: {criteria}")
            
            status, messages = conn.search(None, criteria)
            if status != "OK":
                return [], 0
                
            msg_ids = messages[0].split()
            if not msg_ids:
                return [], 0
                
            # เอาตาม max_results (เอาอันล่าสุดมา)
            # IMAP ids เรียงตามเวลาเก่าไปใหม่
            msg_ids = msg_ids[-WORK_EMAIL_MAX_RESULTS:]
            
            for msg_id in reversed(msg_ids):
                # ดึงเต็มเรื่อง แต่ใช้ BODY.PEEK[] เพื่อไม่ให้กระทบสถานะ \Seen 
                status, msg_data = conn.fetch(msg_id, "(BODY.PEEK[])")
                if status != "OK" or not msg_data[0]:
                    continue
                    
                raw_email = msg_data[0][1]
                msg = email.message_from_bytes(raw_email)
                
                # Header extraction
                subject = self._decode_header(msg.get("Subject", "(ไม่มีหัวข้อ)"))
                sender = self._decode_header(msg.get("From", "(unknown)"))
                date_str = msg.get("Date", "")
                to_str = self._decode_header(msg.get("To", ""))

                # ID Check first before processing intensive things
                msg_id_header = msg.get("Message-ID")
                if msg_id_header and msg_id_header.strip():
                    final_unique_id = msg_id_header.strip()
                else:
                    # Use decoded msg_id to ensure uniqueness natively without Message-ID
                    hash_str = f"{subject}_{sender}_{date_str}_{msg_id.decode('utf-8')}"
                    final_unique_id = "imap_hash_" + hashlib.md5(hash_str.encode('utf-8')).hexdigest()
                    
                if not parsed.force and db.is_email_processed(user_id, final_unique_id):
                    skipped += 1
                    del msg
                    continue
                
                # Body
                body = self._extract_body(msg)
                
                # Python-side filtering for Thai compatibility
                skip_this_msg = False
                for k, v in parsed.filters.items():
                    if k == "folder":
                        continue
                    v_lower = v.lower()
                    if k == "from" and v_lower not in sender.lower(): skip_this_msg = True
                    elif k == "to" and v_lower not in to_str.lower(): skip_this_msg = True
                    elif k == "subject" and v_lower not in subject.lower(): skip_this_msg = True
                    elif k == "body" and v_lower not in body.lower(): skip_this_msg = True
                
                if parsed.search_text:
                    s_lower = parsed.search_text.lower()
                    if s_lower not in subject.lower() and s_lower not in sender.lower() and s_lower not in body.lower():
                        skip_this_msg = True
                        
                if skip_this_msg:
                    del msg
                    continue
                    
                body = body[:2000]
                
                # Attachments
                attachments = self._process_attachments(msg)
                
                emails_data.append(EmailData(
                    message_id=final_unique_id,
                    subject=subject,
                    sender=sender,
                    date=date_str,
                    body=body,
                    attachments=attachments
                ))
                del msg # Free memory
                
        except Exception as e:
            log.error(f"IMAP fetch error: {e}")
            raise e
        finally:
            if conn:
                try:
                    conn.close()
                except Exception:
                    pass
                try:
                    conn.logout()
                except Exception:
                    pass
            
        return emails_data, skipped

    async def execute(self, user_id: str, args: str = "", **kwargs) -> str:
        imap_credentials = self._resolve_imap_credentials(user_id)
        if not imap_credentials:
            return (
                "❌ ยังไม่ได้ตั้งค่า IMAP สำหรับบัญชีของคุณ\n"
                "กรุณาตั้งค่าด้วย /setkey work_imap_host <host>, /setkey work_imap_user <user>, และ /setkey work_imap_password <password>"
            )
            
        parsed = self._parse_args(args)
        
        # Build display label for output
        filters_str = ", ".join([f"{k}:{v}" for k, v in parsed.filters.items()] + ([f"search:{parsed.search_text}"] if parsed.search_text else []))
        display_label = parsed.time_label
        if filters_str:
            display_label += f" ({filters_str})"
            
        try:
            loop = asyncio.get_running_loop()
            
            # 1. Fetch emails (async wrapper)
            # await asyncio.wait_for is used to limit execution time
            task = loop.run_in_executor(None, self._sync_fetch_all, user_id, parsed, imap_credentials)
            emails_data, skipped = await asyncio.wait_for(task, timeout=60.0)
            
            if not emails_data:
                hint = f"ไม่พบอีเมลทำงานใน {display_label}"
                if skipped > 0:
                    hint += f" (ข้ามที่สรุปไปแล้ว {skipped} ฉบับ)"
                    hint += "\nถ้าต้องการสรุปใหม่ทั้งหมดให้เติมคำว่า `force`"
                return hint
                
            # 2. Summarize
            user = get_user_by_id(user_id)
            provider = get_preference(user, "default_llm") if user else "gemini"
            
            MAX_PROMPT_CHARS = 80000
            emails_text = ""
            included_count = 0
            for i, em in enumerate(emails_data, 1):
                email_chunk = f"\n--- Email #{i} ---\n"
                email_chunk += f"From: {em.sender}\n"
                email_chunk += f"Date: {em.date}\n"
                email_chunk += f"Subject: {em.subject}\n"
                email_chunk += f"Content: {em.body}\n"
                if em.attachments:
                    email_chunk += f"Attachments:\n"
                    for j, att in enumerate(em.attachments, 1):
                        size_kb = att.size_bytes // 1024
                        att_info = f"  {j}. [{att.mime_type}] {att.filename} ({size_kb} KB)"
                        if att.content:
                            att_info += f" — extracted:\n     {att.content[:200]}..."
                        else:
                            att_info += f" — status: {att.status}"
                        email_chunk += att_info + "\n"
                else:
                    email_chunk += "Attachments: ไม่มี\n"
                    
                if len(emails_text) + len(email_chunk) > MAX_PROMPT_CHARS:
                    emails_text += "\n\n[... ตัดอีเมลฉบับที่เหลือออกเนื่องจากข้อความยาวเกินไป เพื่อป้องกัน LLM token limit ...]"
                    break
                else:
                    emails_text += email_chunk
                    included_count += 1

            now_str = datetime.now().strftime("%d %b %Y %H:%M")
            system_prompt = (
                f"คุณเป็นผู้ช่วยสรุปอีเมลที่ทำงาน สรุปให้กระชับ เข้าใจง่าย เป็นภาษาไทย\n"
                f"เวลาปัจจุบันที่คุณกำลังอ่านและสรุปคือ: {now_str}\n"
                f"คำแนะนำสำคัญเรื่องเวลา (โปรดทำตาม 2 ขั้นตอนนี้อย่างเคร่งครัด): \n"
                f"1. ทำความเข้าใจเวลาของเหตุการณ์เปรียบเทียบจาก 'วันที่ส่งอีเมล (Date)' เช่น ถ้าย้อนดูอีเมลที่ส่งวันที่ 9 และเนื้อหาบอกว่า 'เมื่อวาน' = เหตุการณ์คือวันที่ 8, 'พรุ่งนี้' = เหตุการณ์คือวันที่ 10, หรือ '10 มี.ค.' = เหตุการณ์คือ 10 มี.ค.\n"
                f"2. เวลาพิมพ์สรุปให้ผู้ใช้ ให้แปลงวันที่จากข้อ 1 มาเทียบกับ 'เวลาปัจจุบัน ({now_str})' เสมอ เช่น ถ้าเหตุการณ์จากข้อ 1 คือวันที่ 10 มี.ค. และเวลาปัจจุบันคือ 10 มี.ค. ต้องใช้คำว่า 'วันนี้' หรือ 'เช้านี้' หรือ '10 มี.ค.' ห้ามสรุปว่า 'พรุ่งนี้' เด็ดขาด!\n\n"
                "รูปแบบการสรุป:\n"
                "- ภาพรวม: สรุปสั้นๆ ว่ามีกี่ฉบับ เรื่องอะไรบ้าง\n"
                "- ต้องดำเนินการ: เมลที่ต้องทำอะไร (ตอบกลับ, อนุมัติ, ตรวจสอบ)\n"
                "- จัดกลุ่มตามประเภท: งาน, การเงิน, HR, IT, อื่นๆ\n"
                "- ไฟล์แนบสำคัญ: สรุปเนื้อหาไฟล์แนบที่สำคัญ (ถ้ามี)\n"
                "- สรุปท้าย: สิ่งที่ควรให้ความสำคัญก่อน\n\n"
                "จัดระเบียบให้มีความชัดเจน ใช้ bullet points และ emoji ให้อ่านง่าย"
            )
            
            prompt_count_text = f"{included_count} ฉบับ (จากทั้งหมด {len(emails_data)} ฉบับ)" if included_count < len(emails_data) else f"{len(emails_data)} ฉบับ"
            resp = await llm_router.chat(
                messages=[{"role": "user", "content": f"สรุปอีเมล {prompt_count_text} ({display_label}):\n{emails_text}"}],
                provider=provider,
                tier=self.preferred_tier,
                system=system_prompt,
            )
            
            # 3. Save to DB
            for em in emails_data:
                db.mark_email_processed(user_id, em.message_id, em.subject, em.sender)
                
            db.log_tool_usage(
                user_id=user_id,
                tool_name=self.name,
                llm_model=resp["model"],
                token_used=resp["token_used"],
                status="success",
                **db.make_log_field("input", f"{len(emails_data)} work emails ({display_label})", kind="work_email_batch_request"),
                **db.make_log_field("output", resp["content"], kind="work_email_summary_text"),
            )
            
            return f"📬 สรุปเมลที่ทำงาน ({display_label}):\n\n{resp['content']}"

        except asyncio.TimeoutError:
            log.error("Work email fetch timeout")
            return "⏳ ใช้เวลานานเกินไป ลองจำกัดช่วงเวลาให้แคบลง เช่น `3d` หรือระบุผู้ส่ง"
        except ssl.SSLError:
            return "❌ SSL error: ตรวจสอบ certificate ของ mail server"
        except imaplib.IMAP4.error as e:
            return f"❌ Login หรือเชื่อมต่อ IMAP ไม่ผ่าน: {e}"
        except Exception as e:
            log.error(f"Work email failed for {user_id}: {e}")
            db.log_tool_usage(
                user_id=user_id,
                tool_name=self.name,
                status="failed",
                **db.make_error_fields(str(e)),
            )
            return f"เกิดข้อผิดพลาด: {e}"

    def get_tool_spec(self) -> dict:
        return {
            "name": self.name,
            "description": (
                "สรุปอีเมลที่ทำงานผ่านระบบ IMAP ขององค์กร. "
                "ใช้เมื่อ user ถามเรื่องอีเมลงาน หรืออีเมลบริษัท. "
                "ไม่ใช่สำหรับ Gmail ส่วนตัว (ใช้ gmail_summary). "
                "เช่น 'มีเมลงานเข้าไหม', 'สรุปเมล บขบ. วันนี้', 'เมลจาก hr'"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "args": {
                        "type": "string",
                        "description": (
                            "ช่วงเวลา (1d, 3d, 7d, 30d), "
                            "ค้นหาตามช่อง (from:{xx}, to:{xx}, subject:{xx}, body:{xx}, folder:{xx}), "
                            "force (บังคับสรุปซ้ำ)"
                        ),
                    }
                },
                "required": [],
            },
        }
